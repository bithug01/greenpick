"""One-off: convert README.md to a Word .docx on the Desktop for easy editing.

Lightweight markdown handling (headings, bullets, tables, code blocks, blockquotes, bold/inline
code, images-as-captions) — enough for this README, not a general markdown engine.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

README = Path(__file__).resolve().parents[1] / "README.md"
OUT = Path.home() / "Desktop" / "GreenPick-README.docx"

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")


def add_inline(paragraph, text: str) -> None:
    """Render **bold**, `code`, *italic* within a line."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1]); run.italic = True
        else:
            paragraph.add_run(part)


def main() -> None:
    lines = README.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.core_properties.title = "GreenPick README"

    i = 0
    pending_table: list[list[str]] = []

    def flush_table():
        nonlocal pending_table
        if not pending_table:
            return
        header, *rows = pending_table
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        for c, cell in enumerate(header):
            p = t.rows[0].cells[c].paragraphs[0]; add_inline(p, cell.strip())
            for r in p.runs: r.bold = True
        for row in rows:
            cells = t.add_row().cells
            for c, cell in enumerate(row[:len(header)]):
                add_inline(cells[c].paragraphs[0], cell.strip())
        doc.add_paragraph()
        pending_table = []

    while i < len(lines):
        line = lines[i]

        # Table rows (| a | b |). Skip the |---|---| separator.
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c for c in line.strip().strip("|").split("|")]
            if set("".join(cells).replace("-", "").replace(":", "").strip()) == set():
                i += 1; continue   # separator row
            pending_table.append(cells)
            i += 1; continue
        else:
            flush_table()

        # Fenced code block
        if line.strip().startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1  # closing fence
            p = doc.add_paragraph()
            run = p.add_run("\n".join(buf)); run.font.name = "Consolas"; run.font.size = Pt(9)
            continue

        # Image -> caption line (Word can't resolve the repo-relative SVG; note it)
        m = re.match(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if m:
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(f"[image: {m.group(1)}  —  {m.group(2)}]")
            r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            i += 1; continue

        # Headings
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=min(level, 4))
            i += 1; continue

        # Blockquote
        if line.strip().startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18)
            add_inline(p, line.strip().lstrip(">").strip())
            i += 1; continue

        # Bullets
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line))
            i += 1; continue

        # Blank / plain paragraph
        if line.strip() == "":
            i += 1; continue
        add_inline(doc.add_paragraph(), line)
        i += 1

    flush_table()
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
